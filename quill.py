#!/usr/bin/env python3
"""
Quill - a sleek word processor with Mendeley Cite-compatible citations.

  * Opens and saves .docx, exports PDF
  * Built-in reference library (BibTeX / RIS / CSL-JSON import, DOI lookup, manual entry)
  * Citations are written into the .docx as Word content controls using the same
    tag layout Mendeley Cite uses, so the Mendeley Cite add-in in Word can recognise them
  * Catppuccin Mocha (dark) and Catppuccin Latte (light) themes

Run:  pip install -r requirements.txt  &&  python quill.py [file.docx]
"""
import base64
import io
import json
import random
import re
import sys
import urllib.parse
import urllib.request
import uuid
from pathlib import Path
from xml.sax.saxutils import quoteattr

from PySide6.QtCore import (QBuffer, QIODevice, QMarginsF, QSettings, QStandardPaths, Qt,
                            QTimer, QUrl, Signal)
from PySide6.QtGui import (QAction, QActionGroup, QColor, QFont, QImage, QKeySequence,
                           QPageLayout, QPageSize, QPdfWriter, QPixmap, QTextBlockFormat,
                           QTextCharFormat, QTextCursor, QTextDocument, QTextImageFormat,
                           QTextListFormat)
from PySide6.QtWidgets import (QApplication, QColorDialog, QComboBox, QDialog,
                               QDialogButtonBox, QDockWidget, QFileDialog,
                               QFontComboBox, QFormLayout, QFrame, QHBoxLayout,
                               QInputDialog, QLabel, QLineEdit, QListWidget,
                               QListWidgetItem, QMainWindow, QMessageBox,
                               QPlainTextEdit, QPushButton, QSizePolicy, QTextEdit, QVBoxLayout,
                               QWidget)

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

APP_NAME = "Quill"
DEFAULT_FONT = "Calibri"
DEFAULT_SIZE = 11
HEADING_SIZES = {1: 20, 2: 16, 3: 13, 4: 12, 5: 11, 6: 11}

# Custom QTextFormat property ids (QTextFormat.UserProperty == 0x100000)
CITE_PROP = 0x100000 + 1   # on a char format: JSON of the citation
BIB_PROP = 0x100000 + 2    # on a block format: block belongs to the bibliography

# Mendeley Cite content-control layout
TAG_PREFIX = "MENDELEY_CITATION_v3_"
BIB_TAG = "MENDELEY_BIBLIOGRAPHY"
CITE_ALIAS = "To edit, see citation options in Mendeley Cite"
BIB_ALIAS = "To edit, see bibliography options in Mendeley Cite"
CSL_SCHEMA = "https://github.com/citation-style-language/schema/raw/master/csl-citation.json"

STYLE_NAMES = ["APA 7", "Harvard", "IEEE", "Vancouver"]
NUMERIC_STYLES = {"IEEE", "Vancouver"}
HEAD_NAMES = {"references", "bibliography", "works cited", "reference list"}

# --------------------------------------------------------------------------- themes
MOCHA = dict(
    name="dark", rosewater="#f5e0dc", pink="#f5c2e7", mauve="#cba6f7", red="#f38ba8",
    peach="#fab387", yellow="#f9e2af", green="#a6e3a1", teal="#94e2d5", blue="#89b4fa",
    lavender="#b4befe", text="#cdd6f4", subtext1="#bac2de", subtext0="#a6adc8",
    overlay0="#6c7086", surface2="#585b70", surface1="#45475a", surface0="#313244",
    base="#1e1e2e", mantle="#181825", crust="#11111b", accent_bg="rgba(203,166,247,0.18)")
LATTE = dict(
    name="light", rosewater="#dc8a78", pink="#ea76cb", mauve="#8839ef", red="#d20f39",
    peach="#fe640b", yellow="#df8e1d", green="#40a02b", teal="#179299", blue="#1e66f5",
    lavender="#7287fd", text="#4c4f69", subtext1="#5c5f77", subtext0="#6c6f85",
    overlay0="#9ca0b0", surface2="#acb0be", surface1="#bcc0cc", surface0="#ccd0da",
    base="#eff1f5", mantle="#e6e9ef", crust="#dce0e8", accent_bg="rgba(136,57,239,0.14)")

MOCHA.update(sapphire="#74c7ec", sky="#89dceb", maroon="#eba0ac", flamingo="#f2cdcd")
LATTE.update(sapphire="#209fb5", sky="#04a5e5", maroon="#e64553", flamingo="#dd7878")

IMG_EXT = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp")
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
MAX_IMG_PX = 780

URL_RE = re.compile(r"(https?://\S+|www\.\S+|\b[\w.+-]+@[\w-]+\.[\w.]+\b|\b10\.\d{4,9}/\S+)")
NUM_RE = re.compile(r"(?<![\w.])\d+(?:[.,]\d+)*%?(?!\w)")
CAPS_RE = re.compile(r"\b[A-Z]{2,}\b")
QUOTE_RE = re.compile(r"“[^”\n]*”|\"[^\"\n]*\"|‘[^’\n]*’")
PAREN_RE = re.compile(r"\([^()\n]*\)")

QSS = """
* { font-family: "Inter", "Segoe UI", "SF Pro Text", "Noto Sans", sans-serif; font-size: 13px; }
QMainWindow, QDialog, QDockWidget { background: @mantle@; color: @text@; }
QWidget { color: @text@; }
QLabel { background: transparent; }
QMenuBar { background: @mantle@; color: @subtext1@; padding: 2px 6px; }
QMenuBar::item { padding: 5px 10px; border-radius: 6px; background: transparent; }
QMenuBar::item:selected { background: @surface0@; color: @text@; }
QMenu { background: @base@; border: 1px solid @surface0@; border-radius: 10px; padding: 6px; }
QMenu::item { padding: 6px 28px 6px 14px; border-radius: 6px; }
QMenu::item:selected { background: @surface0@; }
QMenu::separator { height: 1px; background: @surface0@; margin: 5px 8px; }
QToolBar { background: @mantle@; border: none; padding: 4px 8px; spacing: 3px; }
QToolBar::separator { background: @surface0@; width: 1px; margin: 6px 6px; }
QToolButton { background: transparent; border: none; border-radius: 7px; padding: 5px 9px; color: @subtext1@; }
QToolButton:hover { background: @surface0@; color: @text@; }
QToolButton:checked { background: @accent_bg@; color: @mauve@; }
QToolButton:pressed { background: @surface1@; }
QTextEdit { background: @base@; color: @text@; border: none;
            selection-background-color: @surface2@; selection-color: @text@; }
QLineEdit, QComboBox, QPlainTextEdit { background: @surface0@; border: 1px solid transparent;
            border-radius: 8px; padding: 6px 10px; color: @text@; selection-background-color: @surface2@; }
QLineEdit:focus, QComboBox:focus, QPlainTextEdit:focus { border: 1px solid @mauve@; }
QComboBox::drop-down { border: none; width: 22px; }
QComboBox QAbstractItemView { background: @base@; border: 1px solid @surface0@;
            selection-background-color: @surface0@; selection-color: @text@; outline: 0; }
QPushButton { background: @surface0@; border: none; border-radius: 8px; padding: 7px 14px; color: @text@; }
QPushButton:hover { background: @surface1@; }
QPushButton#primary { background: @mauve@; color: @base@; font-weight: 600; }
QPushButton#primary:hover { background: @lavender@; }
QListWidget { background: @base@; border: 1px solid @surface0@; border-radius: 10px; padding: 4px; outline: 0; }
QListWidget::item { padding: 8px 10px; border-radius: 7px; color: @subtext1@; }
QListWidget::item:hover { background: @mantle@; }
QListWidget::item:selected { background: @accent_bg@; color: @text@; }
QDockWidget::title { background: @mantle@; padding: 10px 12px; font-weight: 600; color: @subtext0@; }
QStatusBar { background: @mantle@; color: @subtext0@; }
QStatusBar::item { border: none; }
QScrollBar:vertical { background: transparent; width: 10px; margin: 2px; }
QScrollBar::handle:vertical { background: @surface1@; border-radius: 4px; min-height: 30px; }
QScrollBar::handle:vertical:hover { background: @surface2@; }
QScrollBar:horizontal { background: transparent; height: 10px; margin: 2px; }
QScrollBar::handle:horizontal { background: @surface1@; border-radius: 4px; min-width: 30px; }
QScrollBar::add-line, QScrollBar::sub-line { height: 0; width: 0; }
QScrollBar::add-page, QScrollBar::sub-page { background: transparent; }
QToolTip { background: @surface0@; color: @text@; border: 1px solid @surface1@; padding: 4px 8px; }
QMessageBox, QInputDialog { background: @base@; }
"""


def build_qss(theme):
    out = QSS
    for k, v in theme.items():
        out = out.replace(f"@{k}@", str(v))
    return out


# --------------------------------------------------------------------------- CSL helpers
def authors_of(it):
    out = []
    for a in it.get("author") or []:
        out.append((a.get("family") or a.get("literal") or "", a.get("given", "")))
    return out


def ini(given, dots=True):
    ls = [p[0].upper() for p in given.replace("-", " ").split() if p]
    return " ".join(l + "." for l in ls) if dots else "".join(ls)


def item_year(it):
    try:
        return str(it["issued"]["date-parts"][0][0])
    except Exception:
        return "n.d."


def author_short(it, amp=" & "):
    a = authors_of(it)
    if not a:
        t = it.get("title", "Untitled")
        return t[:30] + "…" if len(t) > 30 else t
    fams = [x[0] for x in a]
    if len(fams) == 1:
        return fams[0]
    if len(fams) == 2:
        return amp.join(fams)
    return fams[0] + " et al."


def clean(s):
    return (s or "").strip().rstrip(".")


def fix_pages(p, dash="–"):
    return re.sub(r"\s*[-–—]+\s*", dash, p or "")


def display_name(it):
    return f"{author_short(it)} ({item_year(it)}) — {it.get('title', 'Untitled')}"


def cite_text(style, cit_items, numbers):
    items = [ci.get("itemData") or {} for ci in cit_items]
    if style == "APA 7":
        return "(" + "; ".join(f"{author_short(i, ' & ')}, {item_year(i)}" for i in items) + ")"
    if style == "Harvard":
        return "(" + "; ".join(f"{author_short(i, ' and ')} {item_year(i)}" for i in items) + ")"
    nums = sorted(numbers.get(ci.get("id"), 0) for ci in cit_items)
    s = ", ".join(str(n) for n in nums)
    return f"[{s}]" if style == "IEEE" else f"({s})"


def bib_segments(style, it, n):
    """Return [(text, italic), ...] for one bibliography entry."""
    au = authors_of(it)
    title = clean(it.get("title"))
    y = item_year(it)
    cont = clean(it.get("container-title"))
    vol, iss, pg = it.get("volume", ""), it.get("issue", ""), it.get("page", "")
    pub, doi, url = it.get("publisher", ""), it.get("DOI", ""), it.get("URL", "")
    link = f"https://doi.org/{doi}" if doi else url
    seg = []
    if style == "APA 7":
        names = [f"{f}, {ini(g)}".strip().rstrip(",") for f, g in au]
        a = names[0] if len(names) == 1 else (", ".join(names[:-1]) + ", & " + names[-1] if names else "")
        head = (a + " " if a else "") + f"({y}). "
        if cont:
            seg += [(head + title + ". ", False), (cont, True)]
            tail = (f", {vol}" if vol else "") + (f"({iss})" if iss else "") + (f", {fix_pages(pg)}" if pg else "") + "."
            seg.append((tail + (" " + link if link else ""), False))
        else:
            seg += [(head, False), (title, True), (". " + (pub + ". " if pub else "") + link, False)]
    elif style == "Harvard":
        names = [f"{f}, {ini(g)}".strip().rstrip(",") for f, g in au]
        a = names[0] if len(names) == 1 else (", ".join(names[:-1]) + " and " + names[-1] if names else "")
        head = (a + " " if a else "") + f"({y}) "
        if cont:
            seg += [(head + f"‘{title}’, ", False), (cont, True)]
            tail = (f", {vol}" if vol else "") + (f"({iss})" if iss else "") + (f", pp. {fix_pages(pg)}" if pg else "") + "."
            seg.append((tail + (f" doi:{doi}" if doi else (" " + url if url else "")), False))
        else:
            seg += [(head, False), (title, True), (". " + (pub + "." if pub else ""), False)]
    elif style == "IEEE":
        names = [f"{ini(g)} {f}".strip() for f, g in au]
        a = names[0] if len(names) == 1 else (
            " and ".join(names) if len(names) == 2 else (", ".join(names[:-1]) + ", and " + names[-1] if names else ""))
        head = f"[{n}] " + (a + ", " if a else "")
        if cont:
            seg += [(head + f"“{title},” ", False), (cont, True)]
            tail = (f", vol. {vol}" if vol else "") + (f", no. {iss}" if iss else "") + \
                   (f", pp. {fix_pages(pg)}" if pg else "") + f", {y}." + (f" doi: {doi}" if doi else "")
            seg.append((tail, False))
        else:
            seg += [(head, False), (title, True), (f". {pub + ', ' if pub else ''}{y}.", False)]
    else:  # Vancouver
        names = [f"{f} {ini(g, False)}".strip() for f, g in au]
        a = ", ".join(names[:6]) + (", et al" if len(names) > 6 else "")
        s = f"{n}. " + (a + ". " if a else "") + title + ". "
        if cont:
            s += cont + ". " + y + (f";{vol}" if vol else "") + (f"({iss})" if iss else "") + \
                 (f":{fix_pages(pg, '-')}" if pg else "") + "."
        else:
            s += (pub + "; " if pub else "") + y + "."
        seg.append((s, False))
    if seg:
        seg[-1] = (seg[-1][0].rstrip(), seg[-1][1])
    return seg


# --------------------------------------------------------------------------- importers
CSL_TYPES = {"article": "article-journal", "book": "book", "inproceedings": "paper-conference",
             "conference": "paper-conference", "incollection": "chapter", "inbook": "chapter",
             "phdthesis": "thesis", "mastersthesis": "thesis", "techreport": "report",
             "misc": "webpage", "online": "webpage"}


def split_names(s):
    out = []
    for n in re.split(r"\s+and\s+", s or ""):
        n = n.strip()
        if not n:
            continue
        if "," in n:
            f, g = n.split(",", 1)
            out.append({"family": f.strip(), "given": g.strip()})
        else:
            parts = n.split()
            out.append({"family": parts[-1], "given": " ".join(parts[:-1])})
    return out


def parse_bibtex(text):
    items = []
    for m in re.finditer(r"@(\w+)\s*[{(]\s*([^,\s]+)\s*,", text):
        etype = m.group(1).lower()
        if etype in ("comment", "string", "preamble"):
            continue
        i, fields = m.end(), {}
        field_re = re.compile(r"\s*(\w[\w-]*)\s*=\s*")
        while i < len(text):
            fm = field_re.match(text, i)
            if not fm:
                if text[i] in ", \n\r\t":
                    i += 1
                    continue
                break
            name, i = fm.group(1).lower(), fm.end()
            if i >= len(text):
                break
            if text[i] == "{":
                depth, j = 0, i
                while j < len(text):
                    if text[j] == "{":
                        depth += 1
                    elif text[j] == "}":
                        depth -= 1
                        if depth == 0:
                            break
                    j += 1
                val, i = text[i + 1:j], j + 1
            elif text[i] == '"':
                j = text.find('"', i + 1)
                j = len(text) if j < 0 else j
                val, i = text[i + 1:j], j + 1
            else:
                vm = re.compile(r"[^,\n}]+").match(text, i)
                val = vm.group(0).strip() if vm else ""
                i = vm.end() if vm else i + 1
            fields[name] = re.sub(r"[{}]", "", re.sub(r"\s+", " ", val)).strip()
        it = {"id": str(uuid.uuid4()), "type": CSL_TYPES.get(etype, "article-journal"),
              "title": fields.get("title", ""), "author": split_names(fields.get("author", ""))}
        ym = re.search(r"\d{4}", fields.get("year", "") or fields.get("date", ""))
        if ym:
            it["issued"] = {"date-parts": [[int(ym.group(0))]]}
        for src, dst in (("journal", "container-title"), ("booktitle", "container-title"),
                         ("volume", "volume"), ("number", "issue"), ("pages", "page"),
                         ("publisher", "publisher"), ("doi", "DOI"), ("url", "URL")):
            if fields.get(src):
                it[dst] = fields[src]
        items.append(it)
    return items


def parse_ris(text):
    items, cur, authors = [], None, []
    for line in text.splitlines():
        m = re.match(r"^([A-Z][A-Z0-9])  - ?(.*)$", line)
        if not m:
            continue
        tag, val = m.group(1), m.group(2).strip()
        if tag == "TY":
            cur, authors = {"id": str(uuid.uuid4()), "type": "article-journal"}, []
            if val in ("BOOK", "EBOOK"):
                cur["type"] = "book"
            elif val == "CONF":
                cur["type"] = "paper-conference"
            elif val == "CHAP":
                cur["type"] = "chapter"
        elif cur is None:
            continue
        elif tag == "ER":
            cur["author"] = split_names(" and ".join(authors))
            items.append(cur)
            cur = None
        elif tag in ("TI", "T1"):
            cur["title"] = val
        elif tag in ("AU", "A1"):
            authors.append(val)
        elif tag in ("PY", "Y1", "DA"):
            ym = re.search(r"\d{4}", val)
            if ym:
                cur["issued"] = {"date-parts": [[int(ym.group(0))]]}
        elif tag in ("JO", "JF", "T2", "JA"):
            cur["container-title"] = val
        elif tag == "VL":
            cur["volume"] = val
        elif tag == "IS":
            cur["issue"] = val
        elif tag == "SP":
            cur["page"] = val
        elif tag == "EP":
            cur["page"] = cur.get("page", "") + "-" + val
        elif tag == "DO":
            cur["DOI"] = val
        elif tag == "UR":
            cur["URL"] = val
        elif tag == "PB":
            cur["publisher"] = val
    return items


def parse_csl_json(text):
    data = json.loads(text)
    if isinstance(data, dict):
        data = [data]
    out = []
    for it in data:
        it = dict(it)
        it.setdefault("id", str(uuid.uuid4()))
        out.append(it)
    return out


def fetch_doi(doi):
    doi = re.sub(r"^https?://(dx\.)?doi\.org/", "", doi.strip())
    req = urllib.request.Request("https://api.crossref.org/works/" + urllib.parse.quote(doi),
                                 headers={"User-Agent": "Quill/1.0"})
    with urllib.request.urlopen(req, timeout=15) as r:
        m = json.load(r)["message"]
    types = {"journal-article": "article-journal", "book": "book", "proceedings-article": "paper-conference",
             "book-chapter": "chapter", "dissertation": "thesis", "report": "report"}
    issued = m.get("issued") or m.get("published") or {}
    it = {"id": str(uuid.uuid4()), "type": types.get(m.get("type"), "article-journal"),
          "title": (m.get("title") or [""])[0],
          "author": [{"family": a.get("family", a.get("name", "")), "given": a.get("given", "")}
                     for a in m.get("author", [])],
          "container-title": (m.get("container-title") or [""])[0],
          "volume": m.get("volume", ""), "issue": m.get("issue", ""), "page": m.get("page", ""),
          "publisher": m.get("publisher", ""), "DOI": m.get("DOI", doi), "URL": m.get("URL", "")}
    if issued.get("date-parts"):
        it["issued"] = {"date-parts": issued["date-parts"]}
    return {k: v for k, v in it.items() if v not in ("", None)}


# --------------------------------------------------------------------------- library
class Library:
    def __init__(self):
        base = Path(QStandardPaths.writableLocation(QStandardPaths.StandardLocation.AppDataLocation))
        base.mkdir(parents=True, exist_ok=True)
        self.path = base / "library.json"
        self.items = []
        try:
            self.items = json.loads(self.path.read_text("utf-8"))
        except Exception:
            pass

    def save(self):
        try:
            self.path.write_text(json.dumps(self.items, ensure_ascii=False, indent=1), "utf-8")
        except Exception:
            pass

    def find(self, iid):
        return next((i for i in self.items if i.get("id") == iid), None)

    def add(self, item, replace=True):
        item.setdefault("id", str(uuid.uuid4()))
        old = self.find(item["id"])
        if old is not None:
            if replace:
                old.clear()
                old.update(item)
        else:
            self.items.append(item)

    def is_duplicate(self, item):
        doi = (item.get("DOI") or "").lower()
        t = (item.get("title") or "").strip().lower()
        for i in self.items:
            if doi and (i.get("DOI") or "").lower() == doi:
                return True
            if t and (i.get("title") or "").strip().lower() == t and item_year(i) == item_year(item):
                return True
        return False


# --------------------------------------------------------------------------- docx <-> QTextDocument
def default_bf():
    bf = QTextBlockFormat()
    bf.setBottomMargin(8)
    try:
        bf.setLineHeight(118, QTextBlockFormat.LineHeightTypes.ProportionalHeight.value)
    except Exception:
        pass
    return bf


def rid():
    return random.randint(1, 2 ** 31 - 1)


def sdt_tag(sdt):
    pr = sdt.find(qn("w:sdtPr"))
    tag = pr.find(qn("w:tag")) if pr is not None else None
    return tag.get(qn("w:val")) if tag is not None else ""


def decode_citation(tag):
    b = tag[len(TAG_PREFIX):]
    b += "=" * (-len(b) % 4)
    try:
        raw = base64.b64decode(b)
    except Exception:
        try:
            raw = base64.urlsafe_b64decode(b)
        except Exception:
            return None
    try:
        return json.loads(raw.decode("utf-8"))
    except Exception:
        return None


def encode_citation(cit):
    raw = json.dumps(cit, separators=(",", ":"), ensure_ascii=False).encode("utf-8")
    return TAG_PREFIX + base64.b64encode(raw).decode("ascii")


def _apply_run_format(run, fmt, heading):
    f = run.font
    fams = fmt.fontFamilies()
    name = fams[0] if isinstance(fams, (list, tuple)) and fams else (fams if isinstance(fams, str) else None)
    if name and name != DEFAULT_FONT:
        f.name = name
    if fmt.fontPointSize() > 0 and not heading:
        f.size = Pt(fmt.fontPointSize())
    if fmt.fontWeight() >= 600 and not heading:
        f.bold = True
    if fmt.fontItalic():
        f.italic = True
    if fmt.fontUnderline():
        f.underline = True
    if fmt.fontStrikeOut():
        f.strike = True
    if fmt.foreground().style() != Qt.BrushStyle.NoBrush:
        c = fmt.foreground().color()
        f.color.rgb = RGBColor(c.red(), c.green(), c.blue())
    va = fmt.verticalAlignment()
    if va == QTextCharFormat.VerticalAlignment.AlignSuperScript:
        f.superscript = True
    elif va == QTextCharFormat.VerticalAlignment.AlignSubScript:
        f.subscript = True


def _citation_sdt(par, cit_json, text, fmt, heading):
    try:
        cit = json.loads(cit_json)
    except Exception:
        cit = {"citationItems": []}
    cit.setdefault("manualOverride", {"isManuallyOverridden": False, "manualOverrideText": ""})
    cit["manualOverride"]["citeprocText"] = text
    tag = encode_citation(cit)
    sdt = parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtPr><w:alias w:val={quoteattr(CITE_ALIAS)}/>'
        f'<w:tag w:val={quoteattr(tag)}/><w:id w:val="{rid()}"/></w:sdtPr><w:sdtContent/></w:sdt>')
    content = sdt.find(qn("w:sdtContent"))
    r = parse_xml(f'<w:r {nsdecls("w")}/>')
    run = Run(r, par)
    run.text = text
    _apply_run_format(run, fmt, heading)
    content.append(r)
    par._p.append(sdt)


def _resource_image(qdoc, name):
    v = qdoc.resource(QTextDocument.ResourceType.ImageResource, QUrl(name))
    if isinstance(v, QImage):
        return v
    if isinstance(v, QPixmap):
        return v.toImage()
    img = QImage(name)
    return img if not img.isNull() else None


def _write_image(qdoc, par, imf):
    img = _resource_image(qdoc, imf.name())
    if img is None or img.isNull():
        return
    buf = QBuffer()
    buf.open(QIODevice.OpenModeFlag.WriteOnly)
    img.save(buf, "PNG")
    w = imf.width() or img.width()
    h = imf.height() or img.height()
    max_pt = 453.0  # printable width on A4 with 25.4 mm margins
    wp, hp = w * 0.75, h * 0.75
    if wp > max_pt:
        hp, wp = hp * max_pt / wp, max_pt
    par.add_run().add_picture(io.BytesIO(bytes(buf.data())), width=Pt(wp), height=Pt(hp))


def write_docx(qdoc, path):
    d = Document()
    normal = d.styles["Normal"]
    normal.font.name = DEFAULT_FONT
    normal.font.size = Pt(DEFAULT_SIZE)
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Mm(25.4)

    bib = []

    def flush():
        if not bib:
            return
        sdt = parse_xml(
            f'<w:sdt {nsdecls("w")}><w:sdtPr><w:alias w:val={quoteattr(BIB_ALIAS)}/>'
            f'<w:tag w:val="{BIB_TAG}"/><w:id w:val="{rid()}"/></w:sdtPr><w:sdtContent/></w:sdt>')
        bib[0].addprevious(sdt)
        content = sdt.find(qn("w:sdtContent"))
        for p in bib:
            content.append(p)
        bib.clear()

    bullets = (QTextListFormat.Style.ListDisc, QTextListFormat.Style.ListCircle,
               QTextListFormat.Style.ListSquare)
    block = qdoc.begin()
    while block.isValid():
        bf = block.blockFormat()
        is_bib = bf.hasProperty(BIB_PROP)
        if not is_bib:
            flush()
        lvl = bf.headingLevel()
        tl = block.textList()
        if lvl:
            par = d.add_paragraph(style=f"Heading {min(lvl, 4)}")
        elif tl:
            par = d.add_paragraph(style="List Bullet" if tl.format().style() in bullets else "List Number")
        else:
            par = d.add_paragraph()
        al = bf.alignment()
        if al & Qt.AlignmentFlag.AlignHCenter:
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif al & Qt.AlignmentFlag.AlignRight:
            par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif al & Qt.AlignmentFlag.AlignJustify:
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if not tl:
            if bf.leftMargin():
                par.paragraph_format.left_indent = Pt(bf.leftMargin() * 0.75)
            if bf.textIndent():
                par.paragraph_format.first_line_indent = Pt(bf.textIndent() * 0.75)

        it = block.begin()
        while not it.atEnd():
            frag = it.fragment()
            fmt = frag.charFormat()
            if fmt.isImageFormat():
                _write_image(qdoc, par, fmt.toImageFormat())
                it += 1
                continue
            text = frag.text().replace("\ufffc", "")
            if text:
                if fmt.hasProperty(CITE_PROP):
                    _citation_sdt(par, fmt.stringProperty(CITE_PROP), text, fmt, lvl)
                else:
                    run = par.add_run(text.replace("\u2028", "\n"))
                    _apply_run_format(run, fmt, lvl)
            it += 1
        if is_bib:
            bib.append(par._p)
        block = block.next()
    flush()
    d.save(path)


def _list_kind(d, p_el):
    pPr = p_el.pPr
    if pPr is None or pPr.numPr is None or pPr.numPr.numId is None:
        return None
    num_id = pPr.numPr.numId.val
    if num_id == 0:
        return None
    try:
        numbering = d.part.numbering_part.element
        num = numbering.xpath(f'./w:num[@w:numId="{num_id}"]')[0]
        absid = num.xpath("./w:abstractNumId/@w:val")[0]
        fmt = numbering.xpath(
            f'./w:abstractNum[@w:abstractNumId="{absid}"]/w:lvl[@w:ilvl="0"]/w:numFmt/@w:val')[0]
        return "bullet" if fmt == "bullet" else "number"
    except Exception:
        return "bullet"


def read_docx(path, qdoc, lib):
    d = Document(path)
    qdoc.clear()
    qdoc.setDefaultFont(QFont(DEFAULT_FONT, DEFAULT_SIZE))
    cur = QTextCursor(qdoc)
    state = {"first": True, "list": None, "kind": None}

    def start_block(bf):
        if state["first"]:
            cur.setBlockFormat(bf)
            state["first"] = False
        else:
            cur.insertBlock(bf, QTextCharFormat())

    def run_fmt(run, level):
        f = QTextCharFormat()
        if run.font.name:
            f.setFontFamilies([run.font.name])
        if run.font.size:
            f.setFontPointSize(run.font.size.pt)
        elif level:
            f.setFontPointSize(HEADING_SIZES.get(level, 12))
        if run.bold or level:
            f.setFontWeight(700)
        if run.italic:
            f.setFontItalic(True)
        u = run.underline
        if u is True or (u not in (None, False) and "NONE" not in str(u).upper()):
            f.setFontUnderline(True)
        if run.font.strike:
            f.setFontStrikeOut(True)
        try:
            c = run.font.color
            if c is not None and c.rgb is not None and str(c.rgb).upper() != "000000":
                f.setForeground(QColor("#" + str(c.rgb)))
        except Exception:
            pass
        if run.font.superscript:
            f.setVerticalAlignment(QTextCharFormat.VerticalAlignment.AlignSuperScript)
        elif run.font.subscript:
            f.setVerticalAlignment(QTextCharFormat.VerticalAlignment.AlignSubScript)
        return f

    def register(cit):
        for ci in cit.get("citationItems", []):
            data = dict(ci.get("itemData") or {})
            iid = ci.get("id") or data.get("id")
            if not iid:
                continue
            data["id"] = iid
            ci["id"] = iid
            ci["itemData"] = data
            if lib.find(iid) is None:
                lib.add(dict(data), replace=False)

    def insert_run_images(r_el):
        for dr in r_el.iter(qn("w:drawing")):
            blip = dr.find(".//{%s}blip" % A_NS)
            if blip is None:
                continue
            try:
                blob = d.part.related_parts[blip.get(qn("r:embed"))].blob
            except Exception:
                continue
            img = QImage.fromData(blob)
            if img.isNull():
                continue
            w, h = img.width(), img.height()
            ext = dr.find(".//{%s}extent" % WP_NS)
            if ext is not None:
                try:
                    w, h = max(1, int(int(ext.get("cx")) / 9525)), max(1, int(int(ext.get("cy")) / 9525))
                except Exception:
                    pass
            if w > MAX_IMG_PX:
                h, w = int(h * MAX_IMG_PX / w), MAX_IMG_PX
            name = "img_" + uuid.uuid4().hex
            qdoc.addResource(QTextDocument.ResourceType.ImageResource, QUrl(name), img)
            f = QTextImageFormat()
            f.setName(name)
            f.setWidth(w)
            f.setHeight(h)
            cur.insertImage(f)

    def para(el, bib):
        par = Paragraph(el, d._body)
        name = par.style.name if par.style is not None else ""
        level = 0
        m = re.match(r"Heading (\d)", name or "")
        if m:
            level = int(m.group(1))
        elif name == "Title":
            level = 1
        if name and name.startswith("List Bullet"):
            kind = "bullet"
        elif name and name.startswith("List Number"):
            kind = "number"
        else:
            kind = _list_kind(d, el)
        bf = default_bf()
        if level:
            bf.setHeadingLevel(min(level, 6))
            bf.setTopMargin(12)
        a = str(par.alignment).upper() if par.alignment is not None else ""
        if "CENTER" in a:
            bf.setAlignment(Qt.AlignmentFlag.AlignHCenter)
        elif "RIGHT" in a:
            bf.setAlignment(Qt.AlignmentFlag.AlignRight)
        elif "JUSTIFY" in a:
            bf.setAlignment(Qt.AlignmentFlag.AlignJustify)
        if bib:
            bf.setProperty(BIB_PROP, True)
        pf = par.paragraph_format
        if not kind:
            try:
                if pf.left_indent:
                    bf.setLeftMargin(pf.left_indent.pt / 0.75)
                if pf.first_line_indent:
                    bf.setTextIndent(pf.first_line_indent.pt / 0.75)
            except Exception:
                pass
        start_block(bf)
        if kind:
            if state["list"] is None or state["kind"] != kind:
                lf = QTextListFormat()
                lf.setStyle(QTextListFormat.Style.ListDisc if kind == "bullet" else QTextListFormat.Style.ListDecimal)
                state["list"] = cur.createList(lf)
                state["kind"] = kind
            else:
                state["list"].add(cur.block())
        else:
            state["list"] = None

        def add_run(r_el):
            run = Run(r_el, par)
            insert_run_images(r_el)
            text = run.text
            if text:
                cur.insertText(text.replace("\n", "\u2028"), run_fmt(run, level))

        def inline_sdt(sdt):
            tag = sdt_tag(sdt)
            content = sdt.find(qn("w:sdtContent"))
            if content is None:
                return
            runs = content.findall(qn("w:r"))
            if tag.startswith(TAG_PREFIX):
                cit = decode_citation(tag)
                text = "".join(Run(r, par).text for r in runs) or "[citation]"
                fmt = run_fmt(Run(runs[0], par), level) if runs else QTextCharFormat()
                if cit:
                    register(cit)
                    fmt.setProperty(CITE_PROP, json.dumps(cit, separators=(",", ":"), ensure_ascii=False))
                cur.insertText(text, fmt)
            else:
                for r in runs:
                    add_run(r)

        for ch in el:
            if ch.tag == qn("w:r"):
                add_run(ch)
            elif ch.tag in (qn("w:hyperlink"), qn("w:ins"), qn("w:smartTag")):
                for r in ch.findall(qn("w:r")):
                    add_run(r)
            elif ch.tag == qn("w:sdt"):
                inline_sdt(ch)

    def table(el):
        t = Table(el, d._body)
        for row in t.rows:
            start_block(default_bf())
            state["list"] = None
            cur.insertText("\t".join(c.text.replace("\n", " ") for c in row.cells))

    def process(el, bib=False):
        if el.tag == qn("w:p"):
            para(el, bib)
        elif el.tag == qn("w:tbl"):
            table(el)
        elif el.tag == qn("w:sdt"):
            is_bib = sdt_tag(el).startswith(BIB_TAG)
            content = el.find(qn("w:sdtContent"))
            if content is not None:
                for ch in content:
                    process(ch, bib or is_bib)

    for ch in d.element.body:
        if ch.tag != qn("w:sectPr"):
            process(ch)
    lib.save()
    qdoc.clearUndoRedoStacks()
    qdoc.setModified(False)


# --------------------------------------------------------------------------- widgets
class Editor(QTextEdit):
    PAGE_WIDTH = 860
    imageActivated = Signal()

    def insert_image(self, image):
        if image is None or image.isNull():
            return False
        w, h = image.width(), image.height()
        if w > MAX_IMG_PX:
            h, w = int(h * MAX_IMG_PX / w), MAX_IMG_PX
        name = "img_" + uuid.uuid4().hex
        self.document().addResource(QTextDocument.ResourceType.ImageResource, QUrl(name), image)
        f = QTextImageFormat()
        f.setName(name)
        f.setWidth(w)
        f.setHeight(h)
        cur = self.textCursor()
        cur.insertImage(f)
        self.setTextCursor(cur)
        return True

    @staticmethod
    def _image_paths(src):
        if not src.hasUrls():
            return []
        return [u.toLocalFile() for u in src.urls()
                if u.isLocalFile() and u.toLocalFile().lower().endswith(IMG_EXT)]

    def canInsertFromMimeData(self, src):
        return src.hasImage() or bool(self._image_paths(src)) or super().canInsertFromMimeData(src)

    def image_at_cursor(self, cur):
        """Return a cursor selecting an image next to `cur`, or None."""
        for move in (QTextCursor.MoveOperation.Right, QTextCursor.MoveOperation.Left):
            c = QTextCursor(cur)
            c.clearSelection()
            c.movePosition(move, QTextCursor.MoveMode.KeepAnchor)
            if c.hasSelection() and c.charFormat().isImageFormat():
                return c
        return None

    def mouseDoubleClickEvent(self, e):
        c = self.image_at_cursor(self.cursorForPosition(e.pos()))
        if c is not None:
            self.setTextCursor(c)
            self.imageActivated.emit()
            return
        super().mouseDoubleClickEvent(e)

    def __init__(self):
        super().__init__()
        self.setAcceptRichText(True)
        self.setFrameShape(QFrame.Shape.NoFrame)
        self.setTabStopDistance(40)

    def resizeEvent(self, e):
        super().resizeEvent(e)
        m = max(28, (self.width() - self.PAGE_WIDTH) // 2)
        self.setViewportMargins(m, 18, m, 18)

    def insertFromMimeData(self, source):
        paths = self._image_paths(source)
        if paths:
            for pth in paths:
                self.insert_image(QImage(pth))
            return
        if source.hasImage() and not source.hasHtml():
            data = source.imageData()
            img = data if isinstance(data, QImage) else (data.toImage() if hasattr(data, "toImage") else None)
            if self.insert_image(img):
                return
        fmt = self.currentCharFormat()
        if fmt.hasProperty(CITE_PROP):
            fmt.clearProperty(CITE_PROP)
            self.setCurrentCharFormat(fmt)
        super().insertFromMimeData(source)

    def keyPressEvent(self, e):
        if e.key() in (Qt.Key.Key_Return, Qt.Key.Key_Enter):
            was_bib = self.textCursor().blockFormat().hasProperty(BIB_PROP)
            super().keyPressEvent(e)
            if was_bib:
                c = self.textCursor()
                bf = c.blockFormat()
                bf.clearProperty(BIB_PROP)
                c.setBlockFormat(bf)
            return
        t = e.text()
        if t and t.isprintable() and not (e.modifiers() & (Qt.KeyboardModifier.ControlModifier |
                                                            Qt.KeyboardModifier.AltModifier)):
            fmt = self.currentCharFormat()
            if fmt.hasProperty(CITE_PROP):
                fmt.clearProperty(CITE_PROP)
                self.setCurrentCharFormat(fmt)
        super().keyPressEvent(e)


class ItemDialog(QDialog):
    TYPES = [("Journal article", "article-journal"), ("Book", "book"), ("Book chapter", "chapter"),
             ("Conference paper", "paper-conference"), ("Thesis", "thesis"), ("Report", "report"),
             ("Web page", "webpage")]

    def __init__(self, parent=None, item=None):
        super().__init__(parent)
        self.setWindowTitle("Reference")
        self.setMinimumWidth(480)
        self.item = dict(item) if item else {"id": str(uuid.uuid4())}
        it = self.item
        form = QFormLayout()
        form.setSpacing(10)
        self.type = QComboBox()
        for label, val in self.TYPES:
            self.type.addItem(label, val)
        idx = self.type.findData(it.get("type", "article-journal"))
        self.type.setCurrentIndex(max(0, idx))
        self.title = QLineEdit(it.get("title", ""))
        self.authors = QPlainTextEdit()
        self.authors.setPlaceholderText("One per line:  Family, Given")
        self.authors.setFixedHeight(80)
        self.authors.setPlainText("\n".join(f"{f}, {g}".strip(", ") for f, g in authors_of(it)))
        self.year = QLineEdit("" if item_year(it) == "n.d." else item_year(it))
        self.container = QLineEdit(it.get("container-title", ""))
        self.volume = QLineEdit(it.get("volume", ""))
        self.issue = QLineEdit(it.get("issue", ""))
        self.pages = QLineEdit(it.get("page", ""))
        self.publisher = QLineEdit(it.get("publisher", ""))
        self.doi = QLineEdit(it.get("DOI", ""))
        self.url = QLineEdit(it.get("URL", ""))
        for label, w in (("Type", self.type), ("Title", self.title), ("Authors", self.authors),
                         ("Year", self.year), ("Journal / book / site", self.container),
                         ("Volume", self.volume), ("Issue", self.issue), ("Pages", self.pages),
                         ("Publisher", self.publisher), ("DOI", self.doi), ("URL", self.url)):
            form.addRow(label, w)
        box = QDialogButtonBox(QDialogButtonBox.StandardButton.Save | QDialogButtonBox.StandardButton.Cancel)
        box.button(QDialogButtonBox.StandardButton.Save).setObjectName("primary")
        box.accepted.connect(self._ok)
        box.rejected.connect(self.reject)
        lay = QVBoxLayout(self)
        lay.setContentsMargins(18, 18, 18, 18)
        lay.addLayout(form)
        lay.addWidget(box)

    def _ok(self):
        if not self.title.text().strip():
            QMessageBox.warning(self, "Reference", "Please enter a title.")
            return
        it = self.item
        it["type"] = self.type.currentData()
        it["title"] = self.title.text().strip()
        it["author"] = split_names(" and ".join(l.strip() for l in self.authors.toPlainText().splitlines() if l.strip()))
        ym = re.search(r"\d{4}", self.year.text())
        if ym:
            it["issued"] = {"date-parts": [[int(ym.group(0))]]}
        else:
            it.pop("issued", None)
        for key, w in (("container-title", self.container), ("volume", self.volume), ("issue", self.issue),
                       ("page", self.pages), ("publisher", self.publisher), ("DOI", self.doi), ("URL", self.url)):
            v = w.text().strip()
            if v:
                it[key] = v
            else:
                it.pop(key, None)
        self.accept()


# --------------------------------------------------------------------------- main window
class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.settings = QSettings()
        self.lib = Library()
        self.path = None
        self.style_name = self.settings.value("citation_style", "APA 7")
        if self.style_name not in STYLE_NAMES:
            self.style_name = "APA 7"
        self.theme = MOCHA if self.settings.value("theme", "dark") == "dark" else LATTE
        self.aesthetic = str(self.settings.value("aesthetic", "false")).lower() == "true"

        self.editor = Editor()
        self.setCentralWidget(self.editor)
        self.resize(1280, 820)
        self._reset_doc()

        self._build_actions()
        self._build_toolbars()
        self._build_menus()
        self._build_dock()

        self.count_label = QLabel("0 words")
        self.statusBar().addPermanentWidget(self.count_label)

        self.hl_timer = QTimer(self, singleShot=True, interval=150)
        self.hl_timer.timeout.connect(self.update_highlights)
        doc = self.editor.document()
        doc.contentsChanged.connect(self._on_change)
        doc.modificationChanged.connect(lambda _: self._update_title())
        self.editor.currentCharFormatChanged.connect(lambda _: self.sync_format())
        self.editor.cursorPositionChanged.connect(self.sync_format)

        self.apply_theme(self.theme)
        self._update_title()
        self.sync_format()

    # ---- document lifecycle
    def _reset_doc(self):
        doc = self.editor.document()
        doc.clear()
        doc.setDefaultFont(QFont(DEFAULT_FONT, DEFAULT_SIZE))
        cur = QTextCursor(doc)
        cur.setBlockFormat(default_bf())
        doc.clearUndoRedoStacks()
        doc.setModified(False)

    def _update_title(self):
        name = Path(self.path).name if self.path else "Untitled"
        mod = "•  " if self.editor.document().isModified() else ""
        self.setWindowTitle(f"{mod}{name} — {APP_NAME}")

    def _on_change(self):
        self.count_label.setText(f"{len(self.editor.document().toPlainText().split())} words")
        self.hl_timer.start()

    def maybe_save(self):
        if not self.editor.document().isModified():
            return True
        r = QMessageBox.question(self, APP_NAME, "Save changes to this document?",
                                 QMessageBox.StandardButton.Save | QMessageBox.StandardButton.Discard |
                                 QMessageBox.StandardButton.Cancel)
        if r == QMessageBox.StandardButton.Save:
            return self.save()
        return r == QMessageBox.StandardButton.Discard

    def new_doc(self):
        if self.maybe_save():
            self.path = None
            self._reset_doc()
            self._update_title()

    def open_doc(self, path=None):
        if not self.maybe_save():
            return
        if not path:
            path, _ = QFileDialog.getOpenFileName(self, "Open document", "", "Word documents (*.docx)")
        if not path:
            return
        try:
            read_docx(path, self.editor.document(), self.lib)
        except Exception as e:
            QMessageBox.critical(self, APP_NAME, f"Could not open the file:\n{e}")
            return
        self.path = path
        self.refresh_library_list()
        self._update_title()
        self.update_highlights()
        self.statusBar().showMessage(f"Opened {Path(path).name}", 4000)

    def save(self):
        return self.save_as() if not self.path else self._write(self.path)

    def save_as(self):
        path, _ = QFileDialog.getSaveFileName(self, "Save as", self.path or "Untitled.docx",
                                              "Word documents (*.docx)")
        if not path:
            return False
        if not path.lower().endswith(".docx"):
            path += ".docx"
        return self._write(path)

    def _write(self, path):
        try:
            write_docx(self.editor.document(), path)
        except Exception as e:
            QMessageBox.critical(self, APP_NAME, f"Could not save the file:\n{e}")
            return False
        self.path = path
        self.editor.document().setModified(False)
        self._update_title()
        self.statusBar().showMessage(f"Saved {Path(path).name}", 4000)
        return True

    def export_pdf(self):
        default = (Path(self.path).with_suffix(".pdf").name if self.path else "Untitled.pdf")
        path, _ = QFileDialog.getSaveFileName(self, "Export PDF", default, "PDF (*.pdf)")
        if not path:
            return
        if not path.lower().endswith(".pdf"):
            path += ".pdf"
        try:
            w = QPdfWriter(path)
            w.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
            w.setPageMargins(QMarginsF(20, 20, 20, 20), QPageLayout.Unit.Millimeter)
            self.editor.document().clone().print_(w)
        except Exception as e:
            QMessageBox.critical(self, APP_NAME, f"Could not export PDF:\n{e}")
            return
        self.statusBar().showMessage(f"Exported {Path(path).name}", 4000)

    def closeEvent(self, e):
        if self.maybe_save():
            self.settings.setValue("theme", self.theme["name"])
            self.settings.setValue("citation_style", self.style_name)
            e.accept()
        else:
            e.ignore()

    # ---- actions / toolbars / menus
    def _act(self, text, slot, shortcut=None, checkable=False, tip=None):
        a = QAction(text, self)
        a.setCheckable(checkable)
        if shortcut:
            a.setShortcut(QKeySequence(shortcut))
        if checkable:
            a.toggled.connect(slot)
        else:
            a.triggered.connect(lambda _=False, s=slot: s())
        if tip:
            a.setToolTip(tip + (f"  ({shortcut})" if shortcut else ""))
        return a

    def _build_actions(self):
        ed = self.editor
        self.a_new = self._act("New", self.new_doc, "Ctrl+N")
        self.a_open = self._act("Open", self.open_doc, "Ctrl+O", tip="Open a .docx")
        self.a_save = self._act("Save", self.save, "Ctrl+S", tip="Save as .docx")
        self.a_saveas = self._act("Save As…", self.save_as, "Ctrl+Shift+S")
        self.a_pdf = self._act("PDF", self.export_pdf, "Ctrl+E", tip="Export to PDF")
        self.a_undo = self._act("Undo", ed.undo, "Ctrl+Z")
        self.a_redo = self._act("Redo", ed.redo, "Ctrl+Y")
        self.a_find = self._act("Find…", self.find_text, "Ctrl+F")
        self.a_theme = self._act("Theme", self.toggle_theme, "Ctrl+Shift+T", tip="Switch light / dark")

        self.a_bold = self._act("B", lambda c: self.merge_fmt(self._fmt(weight=700 if c else 400)), "Ctrl+B", True, "Bold")
        self.a_italic = self._act("I", lambda c: self.merge_fmt(self._fmt(italic=c)), "Ctrl+I", True, "Italic")
        self.a_under = self._act("U", lambda c: self.merge_fmt(self._fmt(underline=c)), "Ctrl+U", True, "Underline")
        self.a_strike = self._act("S", lambda c: self.merge_fmt(self._fmt(strike=c)), None, True, "Strikethrough")
        self.a_sup = self._act("x²", lambda c: self.merge_fmt(self._fmt(valign="sup" if c else "normal")), None, True, "Superscript")
        self.a_sub = self._act("x₂", lambda c: self.merge_fmt(self._fmt(valign="sub" if c else "normal")), None, True, "Subscript")
        self.a_color = self._act("A", self.pick_color, None, False, "Text colour")

        grp = QActionGroup(self)
        self.a_left = self._act("Left", lambda c: c and ed.setAlignment(Qt.AlignmentFlag.AlignLeft), "Ctrl+L", True)
        self.a_center = self._act("Center", lambda c: c and ed.setAlignment(Qt.AlignmentFlag.AlignHCenter), "Ctrl+Shift+E", True)
        self.a_right = self._act("Right", lambda c: c and ed.setAlignment(Qt.AlignmentFlag.AlignRight), "Ctrl+R", True)
        self.a_just = self._act("Justify", lambda c: c and ed.setAlignment(Qt.AlignmentFlag.AlignJustify), "Ctrl+J", True)
        for a in (self.a_left, self.a_center, self.a_right, self.a_just):
            grp.addAction(a)
        self.a_bullets = self._act("• List", lambda c: self.toggle_list(QTextListFormat.Style.ListDisc), None, True, "Bullet list")
        self.a_numbers = self._act("1. List", lambda c: self.toggle_list(QTextListFormat.Style.ListDecimal), None, True, "Numbered list")
        self.a_indent = self._act("→", lambda: self.indent(1), None, False, "Increase indent")
        self.a_outdent = self._act("←", lambda: self.indent(-1), None, False, "Decrease indent")

        self.a_image = self._act("Image", self.insert_image_dialog, "Ctrl+Shift+I", tip="Insert image")
        self.a_resize = self._act("Resize image…", self.resize_image)
        self.a_aesthetic = self._act("Aesthetic", self.set_aesthetic, "Ctrl+Shift+A", True,
                                     "Syntax-style Catppuccin text colours (view only)")
        self.a_aesthetic.blockSignals(True)
        self.a_aesthetic.setChecked(self.aesthetic)
        self.a_aesthetic.blockSignals(False)

        self.a_cite = self._act("Insert citation", self.insert_citation, "Ctrl+Alt+C")
        self.a_bib = self._act("Insert / update bibliography", self.update_bibliography, "Ctrl+Alt+B")
        self.a_refresh = self._act("Refresh citations", self.refresh_citations)
        self.a_add_ref = self._act("Add reference…", self.add_reference)
        self.a_import = self._act("Import references…", self.import_references)
        self.a_doi = self._act("Add by DOI…", self.lookup_doi)

    def _build_toolbars(self):
        tb = self.addToolBar("File")
        tb.setMovable(False)
        for a in (self.a_new, self.a_open, self.a_save, self.a_pdf):
            tb.addAction(a)
        tb.addSeparator()
        tb.addAction(self.a_undo)
        tb.addAction(self.a_redo)
        tb.addSeparator()
        tb.addAction(self.a_cite)
        tb.addAction(self.a_bib)
        spacer = QWidget()
        spacer.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Preferred)
        tb.addWidget(spacer)
        tb.addAction(self.a_theme)

        self.addToolBarBreak()
        fb = self.addToolBar("Format")
        fb.setMovable(False)
        self.heading = QComboBox()
        self.heading.addItems(["Normal", "Heading 1", "Heading 2", "Heading 3"])
        self.heading.activated.connect(self.set_heading)
        fb.addWidget(self.heading)
        self.font_combo = QFontComboBox()
        self.font_combo.setCurrentFont(QFont(DEFAULT_FONT))
        self.font_combo.setMaximumWidth(190)
        self.font_combo.activated.connect(lambda _: self.merge_fmt(self._fmt(family=self.font_combo.currentFont().family())))
        fb.addWidget(self.font_combo)
        self.size_combo = QComboBox()
        self.size_combo.setEditable(True)
        self.size_combo.addItems([str(s) for s in (8, 9, 10, 11, 12, 14, 16, 18, 20, 24, 28, 36, 48)])
        self.size_combo.setCurrentText(str(DEFAULT_SIZE))
        self.size_combo.setMaximumWidth(72)
        self.size_combo.activated.connect(lambda _: self._apply_size())
        fb.addWidget(self.size_combo)
        fb.addSeparator()
        for a in (self.a_bold, self.a_italic, self.a_under, self.a_strike, self.a_sup, self.a_sub, self.a_color):
            fb.addAction(a)
        bf = QFont(self.font())
        bf.setBold(True)
        fb.widgetForAction(self.a_bold).setFont(bf)
        itf = QFont(self.font())
        itf.setItalic(True)
        fb.widgetForAction(self.a_italic).setFont(itf)
        uf = QFont(self.font())
        uf.setUnderline(True)
        fb.widgetForAction(self.a_under).setFont(uf)
        sf = QFont(self.font())
        sf.setStrikeOut(True)
        fb.widgetForAction(self.a_strike).setFont(sf)
        fb.addSeparator()
        for a in (self.a_left, self.a_center, self.a_right, self.a_just):
            fb.addAction(a)
        fb.addSeparator()
        for a in (self.a_bullets, self.a_numbers, self.a_outdent, self.a_indent):
            fb.addAction(a)
        fb.addSeparator()
        fb.addAction(self.a_image)
        fb.addAction(self.a_aesthetic)
        self.editor.imageActivated.connect(self.resize_image)

    def _build_menus(self):
        mb = self.menuBar()
        m = mb.addMenu("&File")
        for a in (self.a_new, self.a_open, None, self.a_save, self.a_saveas, self.a_pdf):
            m.addSeparator() if a is None else m.addAction(a)
        m.addSeparator()
        m.addAction(self._act("Quit", self.close, "Ctrl+Q"))
        m = mb.addMenu("&Edit")
        for a in (self.a_undo, self.a_redo):
            m.addAction(a)
        m.addSeparator()
        m.addAction(self._act("Cut", self.editor.cut, "Ctrl+X"))
        m.addAction(self._act("Copy", self.editor.copy, "Ctrl+C"))
        m.addAction(self._act("Paste", self.editor.paste, "Ctrl+V"))
        m.addAction(self._act("Select all", self.editor.selectAll, "Ctrl+A"))
        m.addSeparator()
        m.addAction(self.a_find)
        m = mb.addMenu("&Insert")
        m.addAction(self.a_image)
        m.addAction(self.a_resize)
        m = mb.addMenu("&Cite")
        for a in (self.a_cite, self.a_bib, self.a_refresh, None, self.a_add_ref, self.a_import, self.a_doi):
            m.addSeparator() if a is None else m.addAction(a)
        m.addSeparator()
        sm = m.addMenu("Citation style")
        grp = QActionGroup(self)
        for s in STYLE_NAMES:
            a = QAction(s, self, checkable=True)
            a.setChecked(s == self.style_name)
            a.triggered.connect(lambda _=False, s=s: self.set_style(s))
            grp.addAction(a)
            sm.addAction(a)
        self.style_actions = grp
        self.view_menu = mb.addMenu("&View")
        self.view_menu.addAction(self.a_theme)
        self.view_menu.addAction(self.a_aesthetic)

    def _build_dock(self):
        dock = QDockWidget("References", self)
        dock.setFeatures(QDockWidget.DockWidgetFeature.NoDockWidgetFeatures)
        w = QWidget()
        lay = QVBoxLayout(w)
        lay.setContentsMargins(12, 10, 12, 12)
        lay.setSpacing(8)
        self.search = QLineEdit()
        self.search.setPlaceholderText("Search library…")
        self.search.textChanged.connect(self.refresh_library_list)
        lay.addWidget(self.search)
        self.ref_list = QListWidget()
        self.ref_list.setSelectionMode(QListWidget.SelectionMode.ExtendedSelection)
        self.ref_list.setWordWrap(True)
        self.ref_list.itemDoubleClicked.connect(lambda _: self.insert_citation())
        lay.addWidget(self.ref_list, 1)
        b = QPushButton("Insert citation")
        b.setObjectName("primary")
        b.clicked.connect(self.insert_citation)
        lay.addWidget(b)
        row = QHBoxLayout()
        for text, slot in (("Add", self.add_reference), ("Import", self.import_references), ("DOI", self.lookup_doi)):
            btn = QPushButton(text)
            btn.clicked.connect(lambda _=False, s=slot: s())
            row.addWidget(btn)
        lay.addLayout(row)
        row = QHBoxLayout()
        for text, slot in (("Edit", self.edit_reference), ("Delete", self.delete_reference)):
            btn = QPushButton(text)
            btn.clicked.connect(lambda _=False, s=slot: s())
            row.addWidget(btn)
        lay.addLayout(row)
        lay.addSpacing(6)
        lay.addWidget(QLabel("Citation style"))
        self.style_combo = QComboBox()
        self.style_combo.addItems(STYLE_NAMES)
        self.style_combo.setCurrentText(self.style_name)
        self.style_combo.activated.connect(lambda _: self.set_style(self.style_combo.currentText()))
        lay.addWidget(self.style_combo)
        bb = QPushButton("Insert / update bibliography")
        bb.clicked.connect(self.update_bibliography)
        lay.addWidget(bb)
        dock.setWidget(w)
        dock.setMinimumWidth(300)
        self.addDockWidget(Qt.DockWidgetArea.RightDockWidgetArea, dock)
        self.dock = dock
        self.view_menu.addAction(dock.toggleViewAction())
        dock.toggleViewAction().setText("References panel")
        self.refresh_library_list()

    # ---- theme
    def apply_theme(self, theme):
        self.theme = theme
        QApplication.instance().setStyleSheet(build_qss(theme))
        self.settings.setValue("theme", theme["name"])
        self.update_highlights()

    def toggle_theme(self):
        self.apply_theme(LATTE if self.theme is MOCHA else MOCHA)

    # ---- formatting helpers
    def _fmt(self, family=None, size=None, weight=None, italic=None, underline=None, strike=None, valign=None):
        f = QTextCharFormat()
        if family:
            f.setFontFamilies([family])
        if size:
            f.setFontPointSize(size)
        if weight is not None:
            f.setFontWeight(weight)
        if italic is not None:
            f.setFontItalic(italic)
        if underline is not None:
            f.setFontUnderline(underline)
        if strike is not None:
            f.setFontStrikeOut(strike)
        if valign:
            f.setVerticalAlignment({"sup": QTextCharFormat.VerticalAlignment.AlignSuperScript,
                                    "sub": QTextCharFormat.VerticalAlignment.AlignSubScript}.get(
                valign, QTextCharFormat.VerticalAlignment.AlignNormal))
        return f

    def merge_fmt(self, fmt):
        cur = self.editor.textCursor()
        if not cur.hasSelection():
            cur.select(QTextCursor.SelectionType.WordUnderCursor)
        cur.mergeCharFormat(fmt)
        self.editor.mergeCurrentCharFormat(fmt)
        self.editor.setFocus()

    def _apply_size(self):
        try:
            s = float(self.size_combo.currentText())
            if s > 0:
                self.merge_fmt(self._fmt(size=s))
        except ValueError:
            pass

    def pick_color(self):
        c = QColorDialog.getColor(self.editor.textColor(), self, "Text colour")
        if c.isValid():
            f = QTextCharFormat()
            f.setForeground(c)
            self.merge_fmt(f)

    def set_heading(self, idx):
        cur = self.editor.textCursor()
        cur.beginEditBlock()
        bf = cur.blockFormat()
        bf.setHeadingLevel(idx)
        bf.setTopMargin(12 if idx else 0)
        cur.setBlockFormat(bf)
        cur.movePosition(QTextCursor.MoveOperation.StartOfBlock)
        cur.movePosition(QTextCursor.MoveOperation.EndOfBlock, QTextCursor.MoveMode.KeepAnchor)
        f = QTextCharFormat()
        f.setFontPointSize(HEADING_SIZES.get(idx, DEFAULT_SIZE))
        f.setFontWeight(700 if idx else 400)
        cur.mergeCharFormat(f)
        cur.endEditBlock()
        self.editor.mergeCurrentCharFormat(f)
        self.editor.setFocus()

    def toggle_list(self, style):
        cur = self.editor.textCursor()
        cur.beginEditBlock()
        lst = cur.currentList()
        if lst is not None and lst.format().style() == style:
            lst.remove(cur.block())
            bf = cur.blockFormat()
            bf.setIndent(0)
            cur.setBlockFormat(bf)
        else:
            lf = QTextListFormat()
            lf.setStyle(style)
            cur.createList(lf)
        cur.endEditBlock()
        self.sync_format()

    def indent(self, delta):
        cur = self.editor.textCursor()
        lst = cur.currentList()
        if lst is not None:
            lf = lst.format()
            lf.setIndent(max(1, lf.indent() + delta))
            lst.setFormat(lf)
        else:
            bf = cur.blockFormat()
            bf.setLeftMargin(max(0, bf.leftMargin() + 36 * delta))
            cur.setBlockFormat(bf)

    def sync_format(self):
        fmt = self.editor.currentCharFormat()
        cur = self.editor.textCursor()
        pairs = ((self.a_bold, fmt.fontWeight() >= 600), (self.a_italic, fmt.fontItalic()),
                 (self.a_under, fmt.fontUnderline()), (self.a_strike, fmt.fontStrikeOut()),
                 (self.a_sup, fmt.verticalAlignment() == QTextCharFormat.VerticalAlignment.AlignSuperScript),
                 (self.a_sub, fmt.verticalAlignment() == QTextCharFormat.VerticalAlignment.AlignSubScript))
        for a, v in pairs:
            a.blockSignals(True)
            a.setChecked(bool(v))
            a.blockSignals(False)
        al = cur.blockFormat().alignment()
        target = self.a_left
        if al & Qt.AlignmentFlag.AlignHCenter:
            target = self.a_center
        elif al & Qt.AlignmentFlag.AlignRight:
            target = self.a_right
        elif al & Qt.AlignmentFlag.AlignJustify:
            target = self.a_just
        for a in (self.a_left, self.a_center, self.a_right, self.a_just):
            a.blockSignals(True)
            a.setChecked(a is target)
            a.blockSignals(False)
        lst = cur.currentList()
        for a, st in ((self.a_bullets, QTextListFormat.Style.ListDisc), (self.a_numbers, QTextListFormat.Style.ListDecimal)):
            a.blockSignals(True)
            a.setChecked(lst is not None and lst.format().style() == st)
            a.blockSignals(False)
        self.heading.blockSignals(True)
        self.heading.setCurrentIndex(min(3, cur.blockFormat().headingLevel()))
        self.heading.blockSignals(False)
        self.font_combo.blockSignals(True)
        self.font_combo.setCurrentFont(fmt.font())
        self.font_combo.blockSignals(False)
        self.size_combo.blockSignals(True)
        sz = fmt.fontPointSize() or DEFAULT_SIZE
        self.size_combo.setCurrentText(str(int(sz)) if float(sz).is_integer() else str(sz))
        self.size_combo.blockSignals(False)

    def find_text(self):
        t, ok = QInputDialog.getText(self, "Find", "Find text:")
        if ok and t:
            if not self.editor.find(t):
                cur = self.editor.textCursor()
                cur.movePosition(QTextCursor.MoveOperation.Start)
                self.editor.setTextCursor(cur)
                if not self.editor.find(t):
                    self.statusBar().showMessage("Not found", 3000)

    # ---- library panel
    def refresh_library_list(self, *_):
        q = self.search.text().lower() if hasattr(self, "search") else ""
        self.ref_list.clear()
        items = sorted(self.lib.items, key=lambda i: (author_short(i).lower(), item_year(i)))
        for it in items:
            name = display_name(it)
            if q and q not in name.lower() and q not in json.dumps(it, ensure_ascii=False).lower():
                continue
            li = QListWidgetItem(name)
            li.setData(Qt.ItemDataRole.UserRole, it["id"])
            self.ref_list.addItem(li)

    def selected_items(self):
        out = []
        for li in self.ref_list.selectedItems():
            it = self.lib.find(li.data(Qt.ItemDataRole.UserRole))
            if it:
                out.append(it)
        return out

    def add_reference(self):
        dlg = ItemDialog(self)
        if dlg.exec():
            self.lib.add(dlg.item)
            self.lib.save()
            self.refresh_library_list()

    def edit_reference(self):
        sel = self.selected_items()
        if len(sel) != 1:
            self.statusBar().showMessage("Select exactly one reference to edit", 3000)
            return
        dlg = ItemDialog(self, sel[0])
        if dlg.exec():
            self.lib.add(dlg.item)
            self.lib.save()
            self.refresh_library_list()

    def delete_reference(self):
        sel = self.selected_items()
        if not sel:
            return
        if QMessageBox.question(self, APP_NAME, f"Delete {len(sel)} reference(s) from your library?\n"
                                "Citations already in the document are not affected.") == QMessageBox.StandardButton.Yes:
            ids = {i["id"] for i in sel}
            self.lib.items = [i for i in self.lib.items if i["id"] not in ids]
            self.lib.save()
            self.refresh_library_list()

    def import_references(self):
        path, _ = QFileDialog.getOpenFileName(self, "Import references", "",
                                              "References (*.bib *.ris *.json);;All files (*)")
        if not path:
            return
        try:
            text = Path(path).read_text("utf-8", errors="replace")
            ext = Path(path).suffix.lower()
            items = parse_ris(text) if ext == ".ris" else parse_csl_json(text) if ext == ".json" else parse_bibtex(text)
        except Exception as e:
            QMessageBox.critical(self, APP_NAME, f"Could not import:\n{e}")
            return
        added = 0
        for it in items:
            if not self.lib.is_duplicate(it):
                self.lib.add(it)
                added += 1
        self.lib.save()
        self.refresh_library_list()
        self.statusBar().showMessage(f"Imported {added} reference(s) ({len(items) - added} duplicates skipped)", 5000)

    def lookup_doi(self):
        doi, ok = QInputDialog.getText(self, "Add by DOI", "DOI (e.g. 10.1000/xyz123):")
        if not ok or not doi.strip():
            return
        QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
        try:
            it = fetch_doi(doi)
        except Exception as e:
            QApplication.restoreOverrideCursor()
            QMessageBox.warning(self, APP_NAME, f"DOI lookup failed:\n{e}")
            return
        QApplication.restoreOverrideCursor()
        dlg = ItemDialog(self, it)
        if dlg.exec():
            self.lib.add(dlg.item)
            self.lib.save()
            self.refresh_library_list()

    # ---- citations
    def set_style(self, name):
        self.style_name = name
        self.style_combo.setCurrentText(name)
        for a in self.style_actions.actions():
            a.setChecked(a.text() == name)
        self.settings.setValue("citation_style", name)
        self.refresh_citations()

    def collect_spans(self):
        spans = []
        block = self.editor.document().begin()
        while block.isValid():
            it = block.begin()
            while not it.atEnd():
                f = it.fragment()
                fmt = f.charFormat()
                if fmt.hasProperty(CITE_PROP):
                    spans.append((f.position(), f.length(), fmt, fmt.stringProperty(CITE_PROP)))
                it += 1
            block = block.next()
        return spans

    def numbering(self, spans):
        order, items = [], {}
        for _, _, _, js in spans:
            try:
                cit = json.loads(js)
            except Exception:
                continue
            for ci in cit.get("citationItems", []):
                iid = ci.get("id")
                if iid and iid not in items:
                    items[iid] = ci.get("itemData") or {}
                    order.append(iid)
        return {iid: n + 1 for n, iid in enumerate(order)}, order, items

    def refresh_citations(self):
        spans = self.collect_spans()
        if not spans:
            return
        numbers, _, _ = self.numbering(spans)
        cur = QTextCursor(self.editor.document())
        cur.beginEditBlock()
        for pos, length, fmt, js in reversed(spans):
            try:
                cit = json.loads(js)
            except Exception:
                continue
            text = cite_text(self.style_name, cit.get("citationItems", []), numbers)
            cur.setPosition(pos)
            cur.setPosition(pos + length, QTextCursor.MoveMode.KeepAnchor)
            cur.insertText(text, fmt)
        cur.endEditBlock()
        self.update_bibliography(create=False)
        self.update_highlights()

    def insert_citation(self):
        sel = self.selected_items()
        if not sel:
            self.statusBar().showMessage("Select one or more references in the panel first", 4000)
            return
        cit = {"citationID": "MENDELEY_CITATION_" + str(uuid.uuid4()), "properties": {"noteIndex": 0},
               "isEdited": False,
               "manualOverride": {"isManuallyOverridden": False, "citeprocText": "", "manualOverrideText": ""},
               "citationItems": [{"id": i["id"], "itemData": dict(i)} for i in sel],
               "schema": CSL_SCHEMA}
        fmt = self.editor.currentCharFormat()
        fmt.setProperty(CITE_PROP, json.dumps(cit, separators=(",", ":"), ensure_ascii=False))
        cur = self.editor.textCursor()
        cur.insertText("[…]", fmt)
        self.editor.setTextCursor(cur)
        clean_fmt = QTextCharFormat(fmt)
        clean_fmt.clearProperty(CITE_PROP)
        self.editor.setCurrentCharFormat(clean_fmt)
        self.refresh_citations()
        self.editor.setFocus()

    def update_bibliography(self, create=True):
        doc = self.editor.document()
        spans = self.collect_spans()
        numbers, order, items = self.numbering(spans)
        blocks = []
        b = doc.begin()
        while b.isValid():
            if b.blockFormat().hasProperty(BIB_PROP):
                blocks.append(b)
            b = b.next()
        if not blocks and (not create or not spans):
            if create:
                self.statusBar().showMessage("Insert a citation first", 3000)
            return
        if self.style_name in NUMERIC_STYLES:
            entries = [(numbers[i], items[i]) for i in order]
        else:
            ids = sorted(order, key=lambda i: (author_short(items[i]).lower(), item_year(items[i])))
            entries = [(k + 1, items[i]) for k, i in enumerate(ids)]

        if blocks:
            cur = QTextCursor(doc)
            cur.beginEditBlock()
            first, last = blocks[0], blocks[-1]
            start = first.position()
            prev = first.previous()
            if prev.isValid() and prev.blockFormat().headingLevel() > 0 and prev.text().strip().lower() in HEAD_NAMES:
                start = prev.position()
            end = last.position() + last.length() - 1
            cur.setPosition(start)
            cur.setPosition(end, QTextCursor.MoveMode.KeepAnchor)
            cur.removeSelectedText()
        else:
            cur = QTextCursor(self.editor.textCursor())
            cur.beginEditBlock()
            if cur.block().text().strip():
                cur.movePosition(QTextCursor.MoveOperation.EndOfBlock)
                cur.insertBlock(default_bf(), QTextCharFormat())
        hbf = default_bf()
        hbf.setHeadingLevel(1)
        hbf.setTopMargin(14)
        cur.setBlockFormat(hbf)
        hf = QTextCharFormat()
        hf.setFontPointSize(HEADING_SIZES[1])
        hf.setFontWeight(700)
        cur.insertText("References", hf)
        for n, it in entries:
            ebf = default_bf()
            ebf.setProperty(BIB_PROP, True)
            if self.style_name not in NUMERIC_STYLES:
                ebf.setLeftMargin(36)
                ebf.setTextIndent(-36)
            cur.insertBlock(ebf, QTextCharFormat())
            for text, ital in bib_segments(self.style_name, it, n):
                f = QTextCharFormat()
                f.setFontItalic(ital)
                cur.insertText(text, f)
        if not cur.block().next().isValid():
            cur.insertBlock(default_bf(), QTextCharFormat())
        cur.endEditBlock()
        self.update_highlights()

    # ---- images
    def insert_image_dialog(self):
        paths, _ = QFileDialog.getOpenFileNames(self, "Insert image", "",
                                                "Images (*.png *.jpg *.jpeg *.gif *.bmp *.webp)")
        for pth in paths:
            img = QImage(pth)
            if not self.editor.insert_image(img):
                QMessageBox.warning(self, APP_NAME, f"Could not read image:\n{pth}")
        self.editor.setFocus()

    def resize_image(self):
        cur = self.editor.textCursor()
        c = None
        if cur.hasSelection():
            pos = min(cur.anchor(), cur.position())
            probe = QTextCursor(cur)
            probe.setPosition(pos)
            probe.setPosition(pos + 1, QTextCursor.MoveMode.KeepAnchor)
            if probe.charFormat().isImageFormat():
                c = probe
        if c is None:
            c = self.editor.image_at_cursor(cur)
        if c is None:
            self.statusBar().showMessage("Place the cursor next to an image (or double-click it)", 4000)
            return
        imf = c.charFormat().toImageFormat()
        w, h = imf.width() or 1, imf.height() or 1
        nw, ok = QInputDialog.getInt(self, "Resize image", "Width (px):", int(w), 20, 4000)
        if ok:
            imf.setWidth(nw)
            imf.setHeight(max(1, round(h * nw / w)))
            c.setCharFormat(imf)

    # ---- aesthetic (syntax-style colours, display only)
    def set_aesthetic(self, on):
        self.aesthetic = bool(on)
        self.settings.setValue("aesthetic", "true" if on else "false")
        self.update_highlights()

    def _aesthetic_selections(self):
        t, doc, out = self.theme, self.editor.document(), []

        def add(start, length, color):
            if length <= 0:
                return
            c = QTextCursor(doc)
            c.setPosition(start)
            c.setPosition(start + length, QTextCursor.MoveMode.KeepAnchor)
            f = QTextCharFormat()
            f.setForeground(QColor(color))
            sel = QTextEdit.ExtraSelection()
            sel.cursor = c
            sel.format = f
            out.append(sel)

        block = doc.begin()
        while block.isValid() and len(out) < 20000:
            text, base, bf = block.text(), block.position(), block.blockFormat()
            if bf.headingLevel():
                add(base, len(text), t["mauve"])
            else:
                if bf.hasProperty(BIB_PROP):
                    add(base, len(text), t["subtext1"])
                for m in PAREN_RE.finditer(text):
                    add(base + m.start(), m.end() - m.start(), t["sapphire"])
                it = block.begin()
                while not it.atEnd():
                    fr = it.fragment()
                    fm = fr.charFormat()
                    if not fm.hasProperty(CITE_PROP) and fm.foreground().style() == Qt.BrushStyle.NoBrush:
                        if fm.fontWeight() >= 600:
                            add(fr.position(), fr.length(), t["maroon"])
                        elif fm.fontItalic():
                            add(fr.position(), fr.length(), t["pink"])
                    it += 1
                for rx, key in ((NUM_RE, "peach"), (CAPS_RE, "yellow"), (QUOTE_RE, "green"), (URL_RE, "teal")):
                    for m in rx.finditer(text):
                        add(base + m.start(), m.end() - m.start(), t[key])
            block = block.next()
        return out

    def update_highlights(self):
        sels = self._aesthetic_selections() if self.aesthetic else []
        bg = QColor(self.theme["mauve"] if not self.aesthetic else self.theme["blue"])
        bg.setAlpha(55)
        doc = self.editor.document()
        for pos, length, _, _ in self.collect_spans():
            s = QTextEdit.ExtraSelection()
            c = QTextCursor(doc)
            c.setPosition(pos)
            c.setPosition(pos + length, QTextCursor.MoveMode.KeepAnchor)
            f = QTextCharFormat()
            f.setBackground(bg)
            if self.aesthetic:
                f.setForeground(QColor(self.theme["blue"]))
            s.cursor = c
            s.format = f
            sels.append(s)
        self.editor.setExtraSelections(sels)


def main():
    app = QApplication(sys.argv)
    app.setOrganizationName("Quill")
    app.setApplicationName("Quill")
    app.setStyle("Fusion")
    win = MainWindow()
    win.show()
    if len(sys.argv) > 1 and sys.argv[1].lower().endswith(".docx"):
        win.open_doc(sys.argv[1])
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
